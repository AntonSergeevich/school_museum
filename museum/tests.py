import datetime

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from .models import ExcursionRequest, Exhibit, Hall


@override_settings(MEDIA_ROOT="/tmp/museum-test-media")
class SiteTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0)

    def test_pages_open(self):
        ex = Exhibit.published.first()
        hall = Hall.objects.first()
        for url in ["/", "/tour/", "/cockpit/", "/exhibits/", "/about/", "/visit/", "/documents/", "/privacy/",
                    "/consent/", "/accessibility/", "/sitemap.xml", "/robots.txt",
                    ex.get_absolute_url(), hall.get_absolute_url(), "/exhibits/?q=Ил&hall=" + hall.slug]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_cockpit_page_loads_three_locally(self):
        resp = self.client.get("/cockpit/")
        self.assertContains(resp, '"three": "/static/vendor/three/three.module.min.js"')
        self.assertContains(resp, 'type="module"')
        self.assertContains(resp, 'href="/cockpit/"')            # пункт меню
        self.assertContains(self.client.get("/"), 'href="/cockpit/"')

    def test_people_count_plural(self):
        from museum.templatetags.museum_tags import ru_plural
        forms = "имя,имени,имён"
        self.assertEqual([ru_plural(n, forms) for n in (1, 2, 5, 11, 12, 21, 22, 25, 101, 112, 138)],
                         ["имя", "имени", "имён", "имён", "имён", "имя", "имени", "имён", "имя", "имён", "имён"])

    def test_unpublished_exhibit_hidden(self):
        ex = Exhibit.objects.first()
        ex.is_published = False
        ex.save()
        self.assertEqual(self.client.get(ex.get_absolute_url()).status_code, 404)

    def test_fleet_uses_exhibits(self):
        resp = self.client.get("/")
        self.assertContains(resp, 'href="%s"' % Exhibit.objects.get(sky_silhouette="il62").get_absolute_url())

    def test_slug_transliterated(self):
        self.assertEqual(Hall.objects.create(title="Новый зал").slug, "novyy-zal")

    def test_phone_normalized(self):
        from museum.forms import ExcursionRequestForm
        base = {"name": "Иванова", "people": 5, "consent": "on"}
        for raw in ["+7 (912) 345-67-89", "89123456789", "8 912 345 67 89", "+79123456789", "9123456789", "7-912-345-6789"]:
            with self.subTest(raw=raw):
                f = ExcursionRequestForm({**base, "phone": raw})
                self.assertTrue(f.is_valid(), f.errors)
                self.assertEqual(f.cleaned_data["phone"], "+7 (912) 345-67-89")
        for raw in ["+7 (912) 345", "12345", "+7"]:
            with self.subTest(raw=raw):
                self.assertFalse(ExcursionRequestForm({**base, "phone": raw}).is_valid())
        f = ExcursionRequestForm({**base, "phone": "+49 30 1234 5678"})
        self.assertTrue(f.is_valid())
        self.assertEqual(f.cleaned_data["phone"], "+493012345678")

    def test_search_by_several_words(self):
        resp = self.client.get("/people/", {"q": "Мазурук генерал"})
        self.assertContains(resp, "Мазурук")
        self.assertNotContains(resp, "Шкляров Гавриил")
        self.assertContains(resp, 'data-live-region="results"')
        self.assertContains(resp, "search__btn")
        self.assertContains(self.client.get("/exhibits/", {"q": "самописец Ту-154"}), "Бортовой самописец")

    def test_visit_phone_field_has_mask_hooks(self):
        resp = self.client.get("/visit/")
        self.assertContains(resp, "data-phone")
        self.assertContains(resp, 'inputmode="tel"')

    def test_excursion_requires_consent(self):
        data = {"name": "Иванова А. А.", "phone": "+79990000000", "people": 25,
                "desired_date": (datetime.date.today() + datetime.timedelta(days=3)).isoformat()}
        self.client.post(reverse("museum:visit"), data)
        self.assertEqual(ExcursionRequest.objects.count(), 0)
        resp = self.client.post(reverse("museum:visit"), {**data, "consent": "on"})
        self.assertRedirects(resp, reverse("museum:visit"))
        self.assertEqual(ExcursionRequest.objects.count(), 1)

    def test_honeypot(self):
        self.client.post(reverse("museum:visit"), {"name": "bot", "phone": "1", "people": 1,
                                                   "consent": "on", "website": "spam"})
        self.assertEqual(ExcursionRequest.objects.count(), 0)

    def test_markup_mode_only_for_staff(self):
        self.assertNotContains(self.client.get("/tour/?markup=1"), "Режим разметки")
        User.objects.create_superuser("adm", "a@a.ru", "pass-12345-x")
        self.client.login(username="adm", password="pass-12345-x")
        self.assertContains(self.client.get("/tour/?markup=1"), "Режим разметки")

    def test_admin_pages(self):
        User.objects.create_superuser("adm", "a@a.ru", "pass-12345-x")
        self.client.login(username="adm", password="pass-12345-x")
        for m in ["exhibit", "hall", "panorama", "excursionrequest", "document", "timelineevent"]:
            with self.subTest(m=m):
                self.assertEqual(self.client.get(f"/cabinet/museum/{m}/").status_code, 200)
                self.assertEqual(self.client.get(f"/cabinet/museum/{m}/add/").status_code, 200)


import io
import shutil
import tempfile
from pathlib import Path
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from PIL import Image

from config.envfile import parse_env_file
from museum.management.commands.import_drive import clean_title, parse_biography
from museum.models import Document, Panorama, Person
from museum import notify

TMP_MEDIA = tempfile.mkdtemp(prefix="museum-media-")


def image_bytes(size=(4000, 3000), fmt="JPEG", gps=False):
    img = Image.new("RGB", size, (120, 90, 60))
    buf = io.BytesIO()
    kwargs = {}
    if gps:
        exif = Image.Exif()
        exif[0x8825] = {1: "N", 2: (56.0, 1.0, 0.0)}  # GPSInfo
        kwargs["exif"] = exif
    img.save(buf, fmt, **kwargs)
    return buf.getvalue()


class EnvFileTests(SimpleTestCase):
    def test_last_value_wins_quotes_and_brackets(self):
        with tempfile.NamedTemporaryFile("w", suffix=".env", delete=False, encoding="utf-8") as f:
            f.write("# комментарий\nEMAIL_HOST_PASSWORD=\nDEFAULT_FROM_EMAIL=Музей <m@ya.ru>\n"
                    "EMAIL_HOST_PASSWORD='секрет'\nмусор без равно\nexport A=1\n")
        values = parse_env_file(f.name)
        self.assertEqual(values["EMAIL_HOST_PASSWORD"], "секрет")
        self.assertEqual(values["DEFAULT_FROM_EMAIL"], "Музей <m@ya.ru>")
        self.assertEqual(values["A"], "1")


class BiographyParserTests(SimpleTestCase):
    def test_template(self):
        data = parse_biography([
            "Мазурук Илья Павлович", "", "Дата рождения: 7.07.1906", "Место рождения: г. Брест",
            "Звание: генерал-майор", "Награды: два ордена Ленина", "Боевой путь",
            "В армии с 1927 г.", "С 1942 г. — начальник Красноярской воздушной трассы.",
        ], "Мазурук Илья Павлович")
        self.assertEqual(data["birth_date"], "7.07.1906")
        self.assertEqual(data["rank"], "генерал-майор")
        self.assertEqual(data["awards"], "два ордена Ленина")
        self.assertIn("Красноярской воздушной трассы", data["biography"])
        self.assertNotIn("Мазурук", data["biography"])

    def test_clean_title(self):
        self.assertEqual(clean_title("Воспоминания АлСиб 2025-05-30 12_14_09"), "Воспоминания АлСиб")
        self.assertEqual(clean_title("Карты-схемы,  трасса АлСиб"), "Карты-схемы, трасса АлСиб")


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class ImportDriveTests(TestCase):
    def setUp(self):
        import docx
        self.root = Path(tempfile.mkdtemp(prefix="drive-"))
        person = self.root / "Музей" / "Шкляров Гавриил Самуилович"
        person.mkdir(parents=True)
        d = docx.Document()
        for line in ["Шкляров Гавриил Самуилович", "Дата рождения: 14.07.1924",
                     "Звание: младший техник-лейтенант", "Боевой путь", "Старший авиатехник Красноярского аэропорта."]:
            d.add_paragraph(line)
        d.save(person / "Шкляров Гавриил Самуилович.docx")
        (person / "Портрет. 1946 год.png").write_bytes(image_bytes((1600, 2000), "PNG"))
        (person / "Шкляров Г.С..pdf").write_bytes(b"%PDF-1.4 test")
        topic = self.root / "Музей 2" / "Трасса АлСиб"
        topic.mkdir(parents=True)
        for n in ["Воспоминания АлСиб.pdf", "Воспоминания АлСиб 2025-05-30 12_14_09.pdf",
                  "Воспоминания АлСиб 2025-05-30 12_03_45.pdf"]:
            (topic / n).write_bytes(b"%PDF-1.4 test")
        essays = self.root / "Музей 2" / "Сочинения"
        essays.mkdir()
        essay = docx.Document()
        essay.add_paragraph("Мой дедушка")
        essay.add_paragraph("Мой дедушка был лётчиком.")
        essay.save(essays / "Мой_дедушка_2Г.docx")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_publish_and_include_docx(self):
        call_command("import_drive", str(self.root), "--publish", "--include-docx", stdout=io.StringIO())
        p = Person.objects.get()
        self.assertTrue(p.is_published and p.publication_checked)
        self.assertEqual(self.client.get(p.get_absolute_url()).status_code, 200)
        self.assertTrue(Document.objects.filter(topic="Сочинения", is_published=True).exists())
        self.assertEqual(Document.objects.filter(is_published=False).count(), 0)
        self.assertContains(self.client.get("/documents/"), "DOCX")

    def test_publish_rerun_publishes_hidden(self):
        call_command("import_drive", str(self.root), stdout=io.StringIO())
        self.assertFalse(Person.objects.get().is_published)
        call_command("import_drive", str(self.root), "--publish", stdout=io.StringIO())
        self.assertTrue(Person.objects.get().is_published)
        self.assertEqual(Document.objects.filter(is_published=False).count(), 0)
        self.assertEqual(Person.objects.count(), 1)

    def test_table_layout_and_folder_name_docx(self):
        import docx
        folder = self.root / "Музей" / "Авдохин Григорий Егорович"
        folder.mkdir()
        d = docx.Document()
        t = d.add_table(rows=3, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Дата рождения:", "07.10.1913"
        t.cell(1, 0).text, t.cell(1, 1).text = "Звание:", "Капитан"
        t.cell(2, 0).text = "Награды: орден Красной Звезды"
        d.add_paragraph("Боевой путь")
        d.add_paragraph("Техник по вооружению Канской военной авиашколы.")
        d.save(folder / "Авдохин Григорий Егорович.docx")
        (folder / "Портрет. 1945 год.png").write_bytes(image_bytes((400, 500), "PNG"))
        (folder / "Авдохин Г.Е..pdf").write_bytes(b"%PDF-1.4 test")
        only_photo = self.root / "Музей" / "Ермолаев Георгий Михайлович"
        only_photo.mkdir()
        (only_photo / "Портрет.png").write_bytes(image_bytes((400, 500), "PNG"))
        (only_photo / "Ермолаев Г.М..pdf").write_bytes(b"%PDF-1.4 test")
        # так их раньше импортировали — PDF в архив с темой-ФИО
        stale = Document.objects.create(title="Авдохин Г.Е.", topic=folder.name, category="archive",
                                        published_at="2026-10-12", file="documents/a.pdf")

        call_command("import_drive", str(self.root), "--publish", stdout=io.StringIO())
        p = Person.objects.get(full_name="Авдохин Григорий Егорович")
        self.assertEqual(p.birth_date, "07.10.1913")
        self.assertEqual(p.rank, "Капитан")
        self.assertEqual(p.awards, "орден Красной Звезды")
        self.assertIn("Канской", p.biography)
        self.assertTrue(p.portrait and p.source_file)
        self.assertTrue(Person.objects.filter(full_name="Ермолаев Георгий Михайлович").exists())
        self.assertFalse(Document.objects.filter(pk=stale.pk).exists())
        self.assertFalse(Document.objects.filter(topic__in=[folder.name, only_photo.name]).exists())

    def test_topic_folder_not_person_and_broken_docx(self):
        import zipfile
        import docx
        # в корне «Музей» лежит статья, похожая на биографию, — это не человек
        article = docx.Document()
        article.add_paragraph("Из Красноярска в Туруханск летели 36 дней")
        article.add_paragraph("Боевой путь")
        article.save(self.root / "Музей" / "Из Красноярска в Туруханск.docx")
        # биография со ссылкой на картинку, которой нет в архиве (python-docx на таком падает)
        folder = self.root / "Музей" / "Баско Иван Михайлович"
        folder.mkdir()
        src = folder / "tmp.docx"
        d = docx.Document()
        d.add_paragraph("Дата рождения: 1.01.1915")
        img = folder / "pic.png"
        img.write_bytes(image_bytes((50, 50), "PNG"))
        d.add_picture(str(img))
        d.save(src)
        img.unlink()
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(folder / "Баско Иван Михайлович.docx", "w") as zout:
            for item in zin.infolist():
                if not item.filename.startswith("word/media/"):
                    zout.writestr(item, zin.read(item))
        src.unlink()

        call_command("import_drive", str(self.root), "--publish", stdout=io.StringIO(), stderr=io.StringIO())
        self.assertFalse(Person.objects.filter(full_name="Музей").exists())
        self.assertEqual(Person.objects.get(full_name="Баско Иван Михайлович").birth_date, "1.01.1915")

    def test_textbox_paragraphs_read_once(self):
        import docx
        from docx.oxml import parse_xml
        from museum.management.commands.import_drive import docx_paragraphs
        d = docx.Document()
        box = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
               'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
               'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
               'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><mc:AlternateContent>'
               '<mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent>'
               '<w:p><w:r><w:t>Дата рождения: 1.01.1920</w:t></w:r></w:p>'
               '</w:txbxContent></wps:txbx></w:drawing></mc:Choice>'
               '<mc:Fallback><w:pict><v:textbox><w:txbxContent>'
               '<w:p><w:r><w:t>Дата рождения: 1.01.1920</w:t></w:r></w:p>'
               '</w:txbxContent></v:textbox></w:pict></mc:Fallback>'
               '</mc:AlternateContent></w:r></w:p>')
        d.element.body.insert(0, parse_xml(box))
        path = self.root / "box.docx"
        d.save(path)
        self.assertEqual([t for t in docx_paragraphs(path) if t], ["Дата рождения: 1.01.1920"])

    def test_dry_run_writes_nothing(self):
        call_command("import_drive", str(self.root), "--dry-run", stdout=io.StringIO())
        self.assertEqual(Person.objects.count(), 0)
        self.assertEqual(Document.objects.count(), 0)

    def test_import_hidden_and_idempotent(self):
        call_command("import_drive", str(self.root), stdout=io.StringIO())
        p = Person.objects.get()
        self.assertEqual(p.birth_date, "14.07.1924")
        self.assertFalse(p.is_published)
        self.assertTrue(p.portrait.name.endswith(".jpg"))
        self.assertLessEqual(max(Image.open(p.portrait.path).size), 1200)
        self.assertEqual(Document.objects.filter(topic="Трасса АлСиб").count(), 3)
        self.assertFalse(Document.objects.filter(is_published=True).exists())
        self.assertFalse(Document.objects.filter(topic="Сочинения").exists())
        call_command("import_drive", str(self.root), stdout=io.StringIO())
        self.assertEqual(Person.objects.count(), 1)
        self.assertEqual(Document.objects.count(), 3)
        self.assertEqual(self.client.get(p.get_absolute_url()).status_code, 404)
        p.is_published = True
        p.save()
        self.assertContains(self.client.get("/people/"), "Шкляров")


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class ImageTests(TestCase):
    def test_exhibit_photo_resized_and_gps_removed(self):
        ex = Exhibit.objects.create(title="Фото", image=SimpleUploadedFile("p.jpg", image_bytes(gps=True)))
        img = Image.open(ex.image.path)
        self.assertLessEqual(max(img.size), 2400)
        self.assertNotIn(0x8825, img.getexif())

    def test_panorama_must_be_2_to_1(self):
        p = Panorama(title="Не панорама", image=SimpleUploadedFile("x.jpg", image_bytes((3000, 2000))))
        with self.assertRaises(ValidationError):
            p.full_clean()
        ok = Panorama(title="Панорама", image=SimpleUploadedFile("y.jpg", image_bytes((4000, 2000))))
        ok.full_clean()


class NotifyTests(TestCase):
    @override_settings(MAX_BOT_TOKEN="", MUSEUM_NOTIFY_EMAILS=[])
    def test_silent_without_settings(self):
        self.assertFalse(notify.notify_admin("t", "x"))

    @override_settings(MAX_BOT_TOKEN="tok", MAX_USER_ID="42", MAX_CHAT_ID="", MAX_API_URL="https://max.test")
    def test_max_uses_header_not_query(self):
        with mock.patch("urllib.request.urlopen") as op:
            op.return_value.__enter__.return_value.read.return_value = b"{}"
            self.assertTrue(notify.send_max("привет"))
            req = op.call_args[0][0]
        self.assertEqual(req.full_url, "https://max.test/messages?user_id=42")
        self.assertEqual(req.get_header("Authorization"), "tok")
        self.assertNotIn("tok", req.full_url)

    @override_settings(MAX_BOT_TOKEN="tok", MAX_USER_ID="42", MAX_CHAT_ID="")
    def test_max_failure_does_not_lose_request(self):
        with mock.patch("urllib.request.urlopen", side_effect=TimeoutError):
            resp = self.client.post(reverse("museum:visit"), {
                "name": "Петров", "phone": "+7 999 000-00-00", "people": 3, "consent": "on"})
        self.assertRedirects(resp, reverse("museum:visit"))
        self.assertEqual(ExcursionRequest.objects.count(), 1)


class HealthTests(TestCase):
    def test_health(self):
        data = self.client.get("/health/").json()
        self.assertEqual(data["status"], "ok")
        self.assertIn("version", data)


class IconTests(TestCase):
    def test_root_icons_redirect_to_static(self):
        for url in ["/favicon.ico", "/apple-touch-icon.png"]:
            with self.subTest(url=url):
                resp = self.client.get(url)
                self.assertEqual(resp.status_code, 301)
                self.assertIn("/static/img/", resp["Location"])
